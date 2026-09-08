# -*- coding: utf-8 -*-
"""일일보고 엑셀 생성기 — 기존 보고 양식(날짜·구분·작업·진행률·비고·폴더위치) 그대로.

사용법:
    python build_report.py            # 등록된 전체 날짜 생성
    python build_report.py 2026.9.08  # 특정 날짜만

출력: output/일일보고-YYYY-MM-DD.xlsx  (+ .tsv)
근거: 각 저장소 커밋 로그 + HANDOFF.md + Claude Code 세션 기록. 본인(klyhja) 작업만.
"""
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(HERE, "output")
DEV = r"C:\Users\klyhj\dev"


def loc(rel, branch):
    return "{}\\{}\n[{}]".format(DEV, rel, branch)


BE    = loc(r"qmeet\backend2",              "feature/ai_effortEstimate_20260901")
FE    = loc(r"qmeet\front",                 "feature-ai-effortEsttimate_20260901")
AI    = loc(r"qmeet\qmeet_ai",              "feature/orchestrator")
DSGN  = loc(r"qmeet\front_design_prototype", "main")
INFRA = loc(r"qmeet\qmeet-dev-ssh",         "main")
E2E   = loc(r"e2etest\qmeet",               "main")
PWB   = loc(r"e2etest\playwright_base",     "main")
EDU   = loc(r"leadwalk_study\2026\second-half-automation-training\repos",
            "git 아님 · 교육생 5명 폴더")

# 구분 표기 순서 (기존 시트 관례: 백엔드 → 프론트 → ... → 자동화)
REPORTS = {
    "2026.9.08": [
        ("백엔드\n공수산정(85%)", [
            ("AI 추천 테스트 환경(RCMD_PLTF_CD) 저장", "100%", "", BE),
            ("API 문서 공수 기준 정정(환경 1종)", "100%", "", BE),
        ]),
        ("프론트\n공수산정(20%)", [
            ("결과 화면 AI 추천 테스트 환경 기본 선택", "100%", "", FE),
            ("결과 화면 인원·환경 수 즉시 재계산", "100%", "", FE),
            ("9/8 이전 산정 건 기존 값 유지 가드", "100%", "", FE),
            ("결과 배너 공수 숫자 애니메이션", "100%", "", FE),
            ("실제 기획서(11장) 결과 계산식 E2E 검증", "100%",
             "TC 172건·3분 10초·오류 0", FE),
        ]),
        ("AI 서버\n공수산정", [
            ("기획서 기준 테스트 환경 판정 후 큐밋 전달", "100%", "", AI),
            ("판정 환경 1종 기준으로 공수 산정", "100%", "견적 약 3배 과다 산출 수정", AI),
            ("진행 화면 단계 표시 역행 수정", "100%", "9곳 → 0곳", AI),
            ("추천 환경 값 견적 경로 누락 수정", "100%", "실제 AI 재검증 대기", AI),
        ]),
        ("인프라", [
            ("v369 고정 버전 개발환경 구축", "100%",
             "v369.dev.l-walk.com 개통·결함 6건 수정", INFRA),
        ]),
        ("자동화", [
            ("프로젝트 등록 TC 명세 작성(QM101~106)", "80%",
             "QM103 검증 대상 확정 대기", E2E),
            ("프로젝트 등록 자동화 구현(QM101·102·104~106)", "100%",
             "테스트 11건 전체 통과 · 커밋 대기", E2E),
            ("로그인 공통 처리·화면 모듈 2종 구현", "100%",
             "프로젝트명 칸 data-cy 개발팀 요청 필요", E2E),
            ("실패 증적(Evidence) 수집 검증", "100%", "의도적 실패 TC로 확인", E2E),
            ("Base 예제 테스트 정리", "100%", "11건·895줄 삭제, 전체 통과", E2E),
            ("TC 입력 양식 제정", "100%", "양식 2종(문서·엑셀)", PWB + "\n" + E2E),
            ("Q-Meet 데모 프로젝트 영구 위치 이전", "100%", "", E2E),
        ]),
        ("교육", [
            ("하반기 자동화 교육 4주차 과제 피드백", "100%", "5명 전원 전달 완료", EDU),
        ]),
    ],
    "2026.9.07": [
        ("백엔드\n공수산정(85%)", [
            ("진행 단계(substage) 계약 정의·문서 반영", "100%",
             "가짜 AI 서버로 중계 확인", BE),
        ]),
        ("프론트\n공수산정(20%)", [
            ("진행 화면 substage 칩 연동", "100%",
             "AI 서버 2026-09-06 변경 반영", FE),
            ("모바일에서 헤더가 팝업을 덮는 문제 수정", "100%", "리뷰 8번", FE),
            ("실제 기획서(11장) 진행 화면 E2E 완주", "100%", "substage 칩 동작 확인", FE),
            ("프론트 연동 가이드 현황표 갱신", "100%", "", FE),
        ]),
        ("AI 서버\n공수산정", [
            ("기획서 영구 보관 정책·S3 보안 방안 검토", "100%", "정책 검토(코드 변경 없음)", AI),
            ("산정 결과값(totMd·dsgnMd·expDayCnt) 로직 파악", "100%",
             "정책 검토(코드 변경 없음)", AI),
        ]),
        ("시안", [
            ("공수산정 진입 화면 시안 K·L·M·N 재설계", "100%",
             "벤치마크 레이아웃 3종 · 입력→출력 대조", DSGN),
            ("공수산정 PRD·프로젝트 등록 개선 3단계 PRD 작성", "100%", "", DSGN),
            ("한글 파일명에서 프리커밋 훅이 죽던 문제 수정", "100%", "css-guard", DSGN),
            ("폐기 시안 2건·문구변경요청 양식 2건 정리", "100%", "", DSGN),
        ]),
        ("GA4", [
            ("GA4 회원가입 이벤트 연동", "60%",
             "(not set) 원인 규명 · 운영 전용 전송 가드 검증", BE + "\n" + FE),
        ]),
        ("자동화", [
            ("자동화 Base 템플릿 저장소 구축", "100%",
             "껍데기만 유지 · 실제 자동화는 별도 저장소", PWB),
            ("오래된 artifacts 실행 폴더 자동 정리 기능", "100%", "", PWB),
            ("자동화 개발자용 스크립트 작성 가이드 작성", "100%", "", PWB),
            ("Claude 협업 가이드(9장)·규칙 문서 추가", "100%", "", PWB),
            ("브랜치 정리 · Claude 단독 자동화 실증", "100%", "", PWB),
        ]),
    ],
}

FONT      = Font(name="맑은 고딕", size=10)
FONT_PATH = Font(name="Consolas", size=8, color="444444")
_thin = Side(style="thin", color="000000")
BD  = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEF = Alignment(horizontal="left",   vertical="center", wrap_text=True)


def build(date_label, groups):
    wb = Workbook()
    ws = wb.active
    ws.title = "일일보고"

    flat, r = [], 1
    for label, items in groups:
        start = r
        for task, pct, note, path in items:
            ws.cell(r, 2, label if r == start else None)
            ws.cell(r, 3, task)
            ws.cell(r, 4, pct)
            ws.cell(r, 5, note)
            ws.cell(r, 6, path)
            flat.append((date_label if r == 1 else "",
                         label.replace("\n", " ") if r == start else "",
                         task, pct, note, path.replace("\n", " ")))
            r += 1
        if len(items) > 1:
            ws.merge_cells(start_row=start, start_column=2, end_row=r - 1, end_column=2)

    total = r - 1
    ws.cell(1, 1, date_label)
    ws.merge_cells(start_row=1, start_column=1, end_row=total, end_column=1)

    for row in ws.iter_rows(min_row=1, max_row=total, min_col=1, max_col=6):
        for c in row:
            c.border = BD
            c.font = FONT_PATH if c.column == 6 else FONT
            c.alignment = CEN if c.column in (1, 2, 4) else LEF

    for col, width in zip("ABCDEF", (12, 17, 42, 9, 36, 50)):
        ws.column_dimensions[col].width = width
    for i in range(1, total + 1):
        # 폴더 위치가 저장소 2곳(4줄)이면 행 높이를 늘린다
        cell = ws.cell(i, 6).value or ""
        ws.row_dimensions[i].height = 58 if cell.count("\n") >= 3 else 30

    os.makedirs(OUT_DIR, exist_ok=True)
    y, m, d = date_label.split(".")
    stem = "일일보고-{}-{:02d}-{:02d}".format(y, int(m), int(d))
    xlsx = os.path.join(OUT_DIR, stem + ".xlsx")
    tsv = os.path.join(OUT_DIR, stem + ".tsv")

    try:
        wb.save(xlsx)
        saved = xlsx
    except PermissionError:
        saved = os.path.join(OUT_DIR, "_" + stem + ".xlsx")
        wb.save(saved)
        print("  [대기] {} 이 열려 있어 {} 로 저장".format(
            os.path.basename(xlsx), os.path.basename(saved)))

    with open(tsv, "w", encoding="utf-8-sig", newline="") as f:
        for row in flat:
            f.write("\t".join(row) + "\r\n")

    print("  {}  ({}행 x 6열)".format(os.path.relpath(saved, HERE), total))
    print("  {}".format(os.path.relpath(tsv, HERE)))
    return total


def main():
    wanted = sys.argv[1:] or sorted(REPORTS)
    for date_label in wanted:
        if date_label not in REPORTS:
            print("[건너뜀] 등록되지 않은 날짜: {}  (가능: {})".format(
                date_label, ", ".join(sorted(REPORTS))))
            continue
        print("== {} ==".format(date_label))
        build(date_label, REPORTS[date_label])


if __name__ == "__main__":
    main()
