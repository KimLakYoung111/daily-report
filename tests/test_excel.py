# -*- coding: utf-8 -*-
"""엑셀·TSV 렌더 테스트."""
from openpyxl import load_workbook

from report.config import Category, Config
from report.daydata import DayReport, Item, Location, Row
from report.excel import flatten, location_text, write_tsv, write_xlsx

CFG = Config(
    me=("me@example.com",),
    categories=(
        Category(name="백엔드 공수산정", progress="85%", repos=("qmeet/backend2",)),
        Category(name="자동화", repos=("e2etest/qmeet",)),
    ),
    dev_root=r"C:\Users\klyhj\dev",
)

REPORT = DayReport(
    date="2026.9.08",
    rows=(
        Row(
            category="백엔드 공수산정",
            items=(
                Item("추천 환경 저장", "100%", "",
                     (Location("qmeet/backend2", "feature/x"),)),
                Item("문서 기준 정정", "100%", "",
                     (Location("qmeet/backend2", "feature/x"),)),
            ),
        ),
        Row(
            category="자동화",
            items=(
                Item("TC 명세 작성", "80%", "QM103 대기",
                     (Location("e2etest/qmeet", "main"),)),
            ),
        ),
    ),
)


def test_F열에_절대경로와_브랜치를_넣는다():
    got = location_text((Location("qmeet/front", "feature/y"),), CFG.dev_root)
    assert got == "C:\\Users\\klyhj\\dev\\qmeet\\front\n[feature/y]"


def test_저장소가_여러_곳이면_개행으로_이어붙인다():
    got = location_text(
        (Location("a/b", "main"), Location("c/d", "main")), CFG.dev_root
    )
    assert got.count("[main]") == 2
    assert got.count("\n") == 3


def test_날짜는_첫_행에만_구분은_그룹_첫_행에만():
    rows = flatten(REPORT, CFG)
    assert [r[0] for r in rows] == ["2026.9.08", "", ""]
    assert [r[1] for r in rows] == ["백엔드 공수산정(85%)", "", "자동화"]


def test_작업_진행률_비고를_그대로_옮긴다():
    rows = flatten(REPORT, CFG)
    assert rows[2][2] == "TC 명세 작성"
    assert rows[2][3] == "80%"
    assert rows[2][4] == "QM103 대기"


def test_엑셀을_쓰고_행수를_돌려준다(tmp_path):
    path, total = write_xlsx(REPORT, CFG, tmp_path / "out.xlsx")
    assert total == 3
    assert path.exists()


def test_엑셀_열구성과_병합이_양식에_맞는다(tmp_path):
    path, _ = write_xlsx(REPORT, CFG, tmp_path / "out.xlsx")
    ws = load_workbook(path).active
    assert ws.max_row == 3
    assert ws.max_column == 6
    merged = {str(m) for m in ws.merged_cells.ranges}
    assert "A1:A3" in merged      # 날짜 전체 병합
    assert "B1:B2" in merged      # 백엔드 2행 병합
    assert "B3:B3" not in merged  # 1행짜리 구분은 병합하지 않는다
    assert ws.cell(1, 2).value == "백엔드 공수산정(85%)"
    assert ws.cell(3, 4).value == "80%"


def test_엑셀_열너비와_글꼴이_양식에_맞는다(tmp_path):
    path, _ = write_xlsx(REPORT, CFG, tmp_path / "out.xlsx")
    ws = load_workbook(path).active
    widths = [ws.column_dimensions[c].width for c in "ABCDEF"]
    assert widths == [12, 17, 42, 9, 36, 50]
    assert ws.cell(1, 3).font.name == "맑은 고딕"
    assert ws.cell(1, 6).font.name == "Consolas"
    assert ws.cell(1, 6).font.size == 8


def test_대상이_잠겨_있으면_밑줄_접두사로_쓴다(tmp_path, monkeypatch):
    target = tmp_path / "out.xlsx"

    real_save = __import__("openpyxl").Workbook.save
    calls = {"n": 0}

    def flaky(self, filename):
        calls["n"] += 1
        if calls["n"] == 1:
            raise PermissionError("열려 있음")
        return real_save(self, filename)

    monkeypatch.setattr(__import__("openpyxl").Workbook, "save", flaky)
    path, total = write_xlsx(REPORT, CFG, target)
    assert path.name == "_out.xlsx"
    assert total == 3


def test_TSV_를_탭으로_쓴다(tmp_path):
    out = tmp_path / "out.tsv"
    write_tsv(REPORT, CFG, out)
    lines = out.read_text(encoding="utf-8-sig").splitlines()
    assert len(lines) == 3
    assert lines[0].split("\t")[0] == "2026.9.08"
    assert len(lines[0].split("\t")) == 6
    # 셀 안 개행은 한 칸으로 눌러 한 줄을 유지한다
    assert "\n" not in lines[0]


# --- display: name(매칭 키)과 화면 표기를 분리하는 구분 ---
# 사용자 시트는 구분 셀을 "백엔드\n공수산정(85%)" 처럼 두 줄로 보여준다.
# 이 줄바꿈은 이름에서 유도할 수 없는 의미적 구분(지역/모듈)이라
# config.yaml 의 display 필드로 명시한다.

CFG_DISPLAY = Config(
    me=("me@example.com",),
    categories=(
        Category(
            name="백엔드 공수산정",
            display="백엔드\n공수산정",
            progress="85%",
            repos=("qmeet/backend2",),
        ),
    ),
    dev_root=r"C:\Users\klyhj\dev",
)

REPORT_DISPLAY = DayReport(
    date="2026.9.08",
    rows=(
        Row(
            category="백엔드 공수산정",
            items=(
                Item("추천 환경 저장", "100%", "",
                     (Location("qmeet/backend2", "feature/x"),)),
            ),
        ),
    ),
)


def test_display가_있으면_xlsx_B셀에_줄바꿈이_그대로_남는다(tmp_path):
    path, _ = write_xlsx(REPORT_DISPLAY, CFG_DISPLAY, tmp_path / "out.xlsx")
    ws = load_workbook(path).active
    assert "\n" in ws.cell(1, 2).value
    assert ws.cell(1, 2).value == "백엔드\n공수산정(85%)"


def test_TSV는_display의_줄바꿈도_공백으로_누른다(tmp_path):
    out = tmp_path / "out.tsv"
    write_tsv(REPORT_DISPLAY, CFG_DISPLAY, out)
    lines = out.read_text(encoding="utf-8-sig").splitlines()
    # B열(구분) 라벨에 실제 개행이 남아 있으면 TSV 파일의 줄 수 자체가
    # 늘어난다 — F열만 누르고 B열을 그대로 두면 이 assert 가 잡아낸다.
    assert len(lines) == 1
    assert "\n" not in lines[0]
    assert lines[0].split("\t")[1] == "백엔드 공수산정(85%)"
