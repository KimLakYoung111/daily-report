# -*- coding: utf-8 -*-
"""엑셀·TSV 렌더. 기존 보고 양식(6열)을 그대로 유지한다.

열 구성: A 날짜 / B 구분 / C 세부 작업 / D 진행률 / E 비고 / F 폴더 위치
"""
from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from .config import Config
from .daydata import DayReport, Location

WIDTHS = (12, 17, 42, 9, 36, 50)
ROW_HEIGHT = 30
ROW_HEIGHT_TALL = 58

FONT = Font(name="맑은 고딕", size=10)
FONT_PATH = Font(name="Consolas", size=8, color="444444")
_THIN = Side(style="thin", color="000000")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def location_text(locations: Sequence[Location], dev_root: str) -> str:
    """F열 문자열. 위치마다 '<절대경로>\\n[<브랜치>]'."""
    parts = []
    for loc in locations:
        abs_path = str(Path(dev_root) / loc.repo)
        parts.append("{}\n[{}]".format(abs_path, loc.branch))
    return "\n".join(parts)


def flatten(
    report: DayReport, config: Config
) -> list[tuple[str, str, str, str, str, str]]:
    """TSV 용 6열 행 목록. 셀 안 개행은 한 칸으로 눌러 한 줄을 유지한다."""
    out: list[tuple[str, str, str, str, str, str]] = []
    first = True
    for row in report.rows:
        group_first = True
        for item in row.items:
            out.append(
                (
                    report.date if first else "",
                    config.label(row.category) if group_first else "",
                    item.task,
                    item.progress,
                    item.note,
                    location_text(item.locations, config.dev_root).replace("\n", " "),
                )
            )
            first = False
            group_first = False
    return out


def write_xlsx(
    report: DayReport, config: Config, path: str | Path
) -> tuple[Path, int]:
    """엑셀을 쓴다. 대상이 잠겨 있으면 '_' 접두사 경로로 쓴다."""
    wb = Workbook()
    ws = wb.active
    ws.title = "일일보고"

    r = 1
    for row in report.rows:
        start = r
        for item in row.items:
            ws.cell(r, 2, config.label(row.category) if r == start else None)
            ws.cell(r, 3, item.task)
            ws.cell(r, 4, item.progress)
            ws.cell(r, 5, item.note)
            ws.cell(r, 6, location_text(item.locations, config.dev_root))
            r += 1
        if len(row.items) > 1:
            ws.merge_cells(
                start_row=start, start_column=2, end_row=r - 1, end_column=2
            )

    total = r - 1
    ws.cell(1, 1, report.date)
    ws.merge_cells(start_row=1, start_column=1, end_row=total, end_column=1)

    for cells in ws.iter_rows(min_row=1, max_row=total, min_col=1, max_col=6):
        for cell in cells:
            cell.border = BORDER
            cell.font = FONT_PATH if cell.column == 6 else FONT
            cell.alignment = CENTER if cell.column in (1, 2, 4) else LEFT

    for col, width in zip("ABCDEF", WIDTHS):
        ws.column_dimensions[col].width = width
    for i in range(1, total + 1):
        value = ws.cell(i, 6).value or ""
        ws.row_dimensions[i].height = (
            ROW_HEIGHT_TALL if value.count("\n") >= 3 else ROW_HEIGHT
        )

    path = Path(path)
    try:
        wb.save(path)
        return path, total
    except PermissionError:
        fallback = path.with_name("_" + path.name)
        wb.save(fallback)
        return fallback, total


def write_tsv(report: DayReport, config: Config, path: str | Path) -> None:
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        for row in flatten(report, config):
            fh.write("\t".join(row) + "\r\n")
